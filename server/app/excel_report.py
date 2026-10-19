import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .database import as_valid_date, normalize_bill_items, to_number

HEADER_FONT = Font(color='000000', bold=True, size=11)
HEADER_FILL = PatternFill(fill_type='solid', fgColor='FFFFFF')
THIN = Side(style='thin')
BORDER = Border(top=THIN, left=THIN, bottom=THIN, right=THIN)
RS_FORMAT = '"Rs" #,##0.00'


def _style_header(ws, columns: int):
    for col in range(1, columns + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def _style_grid(ws, columns: int, right_aligned_from: int | None = 2, currency_columns=()):
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=columns):
        for index, cell in enumerate(row, start=1):
            cell.border = BORDER
            if right_aligned_from is None:
                horizontal = 'left'
            else:
                horizontal = 'right' if index >= right_aligned_from else 'left'
            cell.alignment = Alignment(vertical='center', horizontal=horizontal)
            if cell.row > 1 and index in currency_columns:
                cell.number_format = RS_FORMAT


def generate_shift_excel_report(report_data: dict, bill_rows, remaining_stock_summary) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = 'Sri Nikil Tradings'
    workbook.properties.created = datetime.now(timezone.utc)

    # Sheet 1: Shift Summary
    summary_sheet = workbook.active
    summary_sheet.title = 'Shift Summary'
    summary_sheet.append(['Field', 'Value'])
    summary_sheet.column_dimensions['A'].width = 30
    summary_sheet.column_dimensions['B'].width = 40
    _style_header(summary_sheet, 2)

    stock_totals = (remaining_stock_summary or {}).get('totals') or {}

    summary_rows = [
        ('Shop Name', 'Sri Nikil Tradings'),
        ('Recipient Email', report_data.get('reportEmail') or 'N/A'),
        ('Shift User', report_data.get('user') or 'N/A'),
        ('Role', report_data.get('role') or 'N/A'),
        ('Shift Start', report_data.get('shiftStartDisplay') or 'N/A'),
        ('Shift End', report_data.get('shiftEndDisplay') or 'N/A'),
        ('Bills Count', report_data.get('billsCount') if report_data.get('billsCount') is not None else 'N/A'),
        ('Total Items Sold', report_data.get('totalItemsSold') if report_data.get('totalItemsSold') is not None else 'N/A'),
        ('Total Sales Amount', f"Rs {to_number(report_data.get('totalSalesAmount')):.2f}"),
        ('Healthy Products', stock_totals.get('healthyCount', 'N/A')),
        ('Low Stock Products', stock_totals.get('lowStockCount', 'N/A')),
        ('Out of Stock Products', stock_totals.get('outOfStockCount', 'N/A')),
    ]
    for row in summary_rows:
        summary_sheet.append(list(row))
    _style_grid(summary_sheet, 2, right_aligned_from=None)

    # Sheet 2: Payment Breakdown
    payment_sheet = workbook.create_sheet('Payment Breakdown')
    payment_sheet.append(['Payment Mode', 'Amount'])
    payment_sheet.column_dimensions['A'].width = 25
    payment_sheet.column_dimensions['B'].width = 25
    _style_header(payment_sheet, 2)

    payment_breakdown = report_data.get('paymentBreakdown') or {}
    total_payments = 0
    for mode in payment_breakdown:
        entry = payment_breakdown[mode]
        amount = (entry or {}).get('amount', 0) if isinstance(entry, dict) else to_number(entry, 0)
        total_payments += amount
        payment_sheet.append([mode, amount])

    payment_sheet.append(['TOTAL', total_payments])
    payment_sheet.cell(row=payment_sheet.max_row, column=1).font = Font(bold=True)
    payment_sheet.cell(row=payment_sheet.max_row, column=2).font = Font(bold=True)
    _style_grid(payment_sheet, 2, right_aligned_from=2, currency_columns=(2,))

    # Sheet 3: Remaining Stock
    stock_sheet = workbook.create_sheet('Remaining Stock')
    stock_headers = [
        'Name of the Product',
        'Quantity Sold',
        'Price Per Product',
        'Total Sale Value Per Product',
        'Opening Stock',
        'Receipts',
        'Sales',
        'Closing Stock',
    ]
    stock_sheet.append(stock_headers)
    for column, width in zip('ABCDEFGH', (42, 18, 20, 28, 18, 15, 15, 18)):
        stock_sheet.column_dimensions[column].width = width
    _style_header(stock_sheet, 8)

    for item in (remaining_stock_summary or {}).get('products') or []:
        opening_stock = to_number(item.get('openingStock', item.get('estimatedOpeningStock', 0)))
        sold_in_shift = to_number(item.get('soldInShift', item.get('sold', 0)))
        refilled_in_shift = to_number(item.get('refilledInShift', item.get('refilled', 0)))
        price = to_number(item.get('price'))
        current_stock = to_number(item.get('currentStock'), opening_stock - sold_in_shift + refilled_in_shift)

        stock_sheet.append([
            item.get('name') or 'N/A',
            sold_in_shift,
            price,
            sold_in_shift * price,
            opening_stock,
            refilled_in_shift,
            sold_in_shift,
            current_stock,
        ])
    _style_grid(stock_sheet, 8, right_aligned_from=2, currency_columns=(3, 4))

    # Sheet 4: Sales Details
    sales_sheet = workbook.create_sheet('Sales Details')
    sales_headers = [
        'Bill No',
        'Date',
        'Customer',
        'Phone',
        'Payment',
        'Items',
        'Subtotal',
        'CGST',
        'SGST',
        'Grand Total',
        'Issued By',
    ]
    sales_sheet.append(sales_headers)
    for column, width in zip('ABCDEFGHIJK', (15, 22, 25, 15, 15, 45, 15, 12, 12, 15, 15)):
        sales_sheet.column_dimensions[column].width = width
    _style_header(sales_sheet, 11)

    for bill in bill_rows or []:
        parsed_items = normalize_bill_items(bill.get('items'))
        date_str = as_valid_date(bill.get('date')).astimezone().strftime('%d/%m/%Y, %H:%M')
        items_list = ', '.join(f"{item.get('name')} x{item.get('qty')}" for item in parsed_items)

        sales_sheet.append([
            bill.get('billNo') or 'N/A',
            date_str,
            bill.get('customer') or 'N/A',
            bill.get('phone') or 'N/A',
            bill.get('payment') or 'N/A',
            items_list or 'No Items',
            to_number(bill.get('subtotal')),
            to_number(bill.get('cgst')),
            to_number(bill.get('sgst')),
            to_number(bill.get('grand')),
            bill.get('by_user') or report_data.get('user') or 'N/A',
        ])
    _style_grid(sales_sheet, 11, right_aligned_from=None, currency_columns=(7, 8, 9, 10))

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
